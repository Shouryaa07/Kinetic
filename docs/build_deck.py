"""Builds docs/Kinetic-Presentation.pptx. Run: python docs/build_deck.py (needs python-pptx)."""
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "screenshots")

NAVY = RGBColor(0x0D, 0x0F, 0x2B)
INK = RGBColor(0x0F, 0x12, 0x26)
MUTED = RGBColor(0x63, 0x6A, 0x86)
BRAND = RGBColor(0x5B, 0x4D, 0xF5)
BRAND_SOFT = RGBColor(0xEC, 0xEA, 0xFE)
CYAN = RGBColor(0x08, 0x91, 0xB2)
LIME = RGBColor(0xC6, 0xF4, 0x32)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
LINE = RGBColor(0xE5, 0xE7, 0xF1)
BG = RGBColor(0xF6, 0xF7, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
page = {"n": 0}


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, content, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, spacing=None):
    """content: str or list of paragraphs; a paragraph is str or list of (text, {opts}) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing:
            para.space_after = Pt(spacing)
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def label_box(slide, x, y, w, h, title, sub=None, fill=WHITE, line=LINE, tcolor=INK, scolor=MUTED,
              tsize=14, ssize=11, radius=0.12):
    box(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    paras = [[(title, {"bold": True, "size": tsize, "color": tcolor})]]
    if sub:
        paras.append([(sub, {"size": ssize, "color": scolor})])
    text(slide, x + 0.05, y, w - 0.1, h, paras, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.75, dash=False, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
        t.set("w", "med")
        t.set("len", "med")
    return c


def new_slide(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    page["n"] += 1
    set_bg(s, BG)
    box(s, 0, 0, 13.333, 0.12, fill=BRAND, line=None, shape=MSO_SHAPE.RECTANGLE)
    if kicker:
        text(s, 0.6, 0.38, 9, 0.3, kicker.upper(), size=11, color=BRAND, bold=True)
    text(s, 0.6, 0.62, 12, 0.7, title, size=30, color=INK, bold=True)
    text(s, 0.6, 7.05, 6, 0.3, "Kinetic  |  Java Servlets + JDBC", size=10, color=MUTED)
    text(s, 12.0, 7.05, 0.8, 0.3, str(page["n"]), size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    return s


def bullets(slide, x, y, w, items, size=16, gap=9, color=INK, bullet_color=BRAND):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            paras.append([("●  ", {"color": bullet_color, "size": size - 4}), (it[0], {"bold": True, "size": size}),
                          (" " + it[1], {"size": size, "color": MUTED})])
        else:
            paras.append([("●  ", {"color": bullet_color, "size": size - 4}), (it, {"size": size})])
    return text(slide, x, y, w, 5, paras, size=size, color=color, spacing=gap)


def code(slide, x, y, w, h, lines, size=11):
    box(slide, x, y, w, h, fill=NAVY, line=None, radius=0.04)
    paras = []
    for ln in lines:
        paras.append([(ln, {"font": "Courier New", "size": size, "color": RGBColor(0xE6, 0xE8, 0xFF)})])
    text(slide, x + 0.2, y + 0.15, w - 0.4, h - 0.3, paras, size=size)


def picture(slide, name, x, y, w=None, h=None, frame=True):
    path = os.path.join(SHOTS, name)
    from PIL import Image
    iw, ih = Image.open(path).size
    if w is not None and h is None:
        h = w * ih / iw
    elif h is not None and w is None:
        w = h * iw / ih
    if frame:
        box(slide, x - 0.06, y - 0.06, w + 0.12, h + 0.12, fill=WHITE, line=LINE, radius=0.02)
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(BLANK)
page["n"] += 1
set_bg(s, NAVY)
box(s, 8.3, -1.5, 7, 7, fill=RGBColor(0x1F, 0x1B, 0x6B), line=None, shape=MSO_SHAPE.OVAL)
box(s, 10.2, 3.6, 5, 5, fill=RGBColor(0x0E, 0x4F, 0x6B), line=None, shape=MSO_SHAPE.OVAL)
box(s, 0.9, 1.1, 0.9, 0.9, fill=BRAND, line=None, radius=0.25)
text(s, 0.9, 1.1, 0.9, 0.9, "K", size=40, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 2.3, 8, 1.2, "Kinetic", size=66, color=WHITE, bold=True)
text(s, 0.9, 3.6, 7.5, 1.2, "A fitness tracking web application with workouts, weekly goals, streaks and Google sign-in",
     size=22, color=RGBColor(0xC9, 0xCD, 0xF5))
box(s, 0.9, 5.2, 0.12, 0.12, fill=LIME, line=None, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.15, 5.08, 8, 0.4, "Java Web Project  |  Servlets, JDBC, core Java", size=16, color=LIME, bold=True)
picture(s, "dashboard.png", 8.1, 1.7, w=4.7)

# ---------------------------------------------------------------- 2. Problem and solution
s = new_slide("The problem and our solution", "Problem understanding")
label_box(s, 0.6, 1.7, 5.9, 0.6, "The problem", fill=ORANGE, line=None, tcolor=WHITE, tsize=16)
bullets(s, 0.7, 2.5, 5.7, [
    "Workout data is scattered across notes, watches and spreadsheets",
    "People cannot tell if they are improving week to week",
    "Weekly targets are set but never checked",
    "Streaks and personal bests are invisible",
], size=17)
label_box(s, 6.9, 1.7, 5.85, 0.6, "Our solution: Kinetic", fill=BRAND, line=None, tcolor=WHITE, tsize=16)
bullets(s, 7.0, 2.5, 5.7, [
    "One log for 7 workout types, with calories estimated from body weight",
    "Dashboard with weekly totals, trends and time-by-type",
    "Goal rings for workouts, minutes, calories and distance",
    "Streaks, activity calendar and personal records",
    "Fast sign-in with Google or email",
], size=17)
box(s, 0.6, 5.5, 12.15, 1.1, fill=BRAND_SOFT, line=None)
text(s, 0.9, 5.5, 11.6, 1.1,
     [[("Who it is for:  ", {"bold": True, "color": BRAND}),
       ("anyone who trains regularly and wants clear numbers, not a complicated app. Built as a full Servlet and JDBC web application with no framework.", {})]],
     size=16, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 3. Features
s = new_slide("Key features", "What it does")
cards = [
    ("Log and edit workouts", "Run, walk, cycle, swim, strength, HIIT, yoga. Duration, distance, effort, notes."),
    ("Smart calories", "MET x weight x hours. Each activity class supplies its own MET value."),
    ("Dashboard", "Today rings, week vs last week, 12-week chart, goal rings, time-by-type donut."),
    ("Weekly goals", "Targets for sessions, minutes, calories, distance. Saved in one transaction."),
    ("Streaks and records", "Current and best streak, longest distance, biggest session."),
    ("Google sign-in", "Continue with Google. Token verified on the server before a session starts."),
    ("Daily health", "Steps, water and sleep with daily rings, plus a body-weight trend line."),
    ("Install on a phone", "Installable web app (PWA), CSV export, light and dark mode, 17-week calendar."),
]
for i, (t, d) in enumerate(cards):
    cx = 0.6 + (i % 4) * 3.07
    cy = 1.7 + (i // 4) * 2.55
    box(s, cx, cy, 2.9, 2.3, fill=WHITE, line=LINE, radius=0.06)
    box(s, cx + 0.2, cy + 0.22, 0.45, 0.45, fill=BRAND_SOFT, line=None, radius=0.3)
    text(s, cx + 0.2, cy + 0.22, 0.45, 0.45, str(i + 1), size=14, color=BRAND, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, cx + 0.2, cy + 0.8, 2.55, 0.4, t, size=16, bold=True)
    text(s, cx + 0.2, cy + 1.2, 2.55, 1.1, d, size=12, color=MUTED)

# ---------------------------------------------------------------- 4. Architecture
s = new_slide("Architecture", "Solution design")
names = [("Browser", "HTML, CSS, JS"), ("Filters", "Security, Auth"), ("Servlets", "15, extend BaseServlet"),
         ("Services", "Business rules"), ("DAOs", "One per table"), ("Database", "Pool + JDBC")]
fills = [WHITE, BRAND_SOFT, BRAND, BRAND, CYAN, NAVY]
for i, (t, sub) in enumerate(names):
    x = 0.6 + i * 2.1
    dark = fills[i] in (BRAND, CYAN, NAVY)
    label_box(s, x, 1.8, 1.8, 1.1, t, sub, fill=fills[i], line=LINE if not dark else None,
              tcolor=WHITE if dark else INK, scolor=RGBColor(0xDD, 0xDF, 0xFF) if dark else MUTED, tsize=15)
    if i < 5:
        arrow(s, x + 1.82, 2.35, x + 2.08, 2.35, color=MUTED, width=2)
label_box(s, 4.7, 3.7, 1.8, 0.9, "Views", "HTML and SVG charts", fill=WHITE)
arrow(s, 5.6, 3.7, 5.6, 2.92, color=BRAND, dash=True)
label_box(s, 6.9, 3.7, 1.8, 0.9, "AuditLogger", "worker thread", fill=WHITE)
arrow(s, 7.8, 3.7, 7.8, 2.92, color=BRAND, dash=True)
label_box(s, 9.1, 3.7, 1.8, 0.9, "Google verifier", "tokeninfo check", fill=WHITE)
arrow(s, 10.0, 3.7, 10.0, 2.92, color=BRAND, dash=True)
text(s, 0.6, 3.75, 3.8, 1.0, [[("Layered design: ", {"bold": True, "color": BRAND}),
                              ("each layer only talks to the one below it.", {})]], size=14)
box(s, 0.6, 5.05, 12.15, 1.25, fill=WHITE, line=LINE, radius=0.05)
bullets(s, 0.85, 5.2, 5.8, [("Composition root:", "AppContext wires every DAO and service"),
                            ("Thin servlets:", "validate, call a service, redirect")], size=14, gap=6)
bullets(s, 6.9, 5.2, 5.7, [("Cross-cutting:", "SecurityFilter (CSP, CSRF), AuthFilter"),
                           ("Observer:", "StatsService cache cleared on every data change")], size=14, gap=6)

# ---------------------------------------------------------------- 5. Tech stack and structure
s = new_slide("Technology and project structure", "Setup")
techs = [("Java 17", "records, switch expressions"), ("Jakarta Servlet 5", "servlets, filters, listeners"),
         ("JDBC + H2", "embedded, swappable to MySQL"), ("Maven + Jetty", "mvn jetty:run"),
         ("JUnit 4", "47 tests"), ("Google Identity", "Sign in with Google")]
for i, (t, d) in enumerate(techs):
    cx = 0.6 + (i % 2) * 3.0
    cy = 1.7 + (i // 2) * 1.45
    box(s, cx, cy, 2.85, 1.25, fill=WHITE, line=LINE, radius=0.08)
    text(s, cx + 0.2, cy + 0.2, 2.5, 0.4, t, size=16, bold=True, color=BRAND)
    text(s, cx + 0.2, cy + 0.65, 2.5, 0.5, d, size=12, color=MUTED)
code(s, 6.9, 1.7, 5.85, 4.85, [
    "kinetic/",
    "  pom.xml, README.md",
    "  src/main/java/com/kinetic/",
    "    model/      Activity hierarchy, enums",
    "    exception/  KineticException tree",
    "    db/         ConnectionPool, Database",
    "    dao/        Dao<T,ID>, JdbcDao, 6 DAOs",
    "    service/    Auth, Activity, Goal, Stats",
    "    concurrent/ AuditLogger, OnlineUsers",
    "    util/       PasswordHasher, LruCache...",
    "    web/        filters, servlet/, view/",
    "  src/main/resources/",
    "    schema.sql, application.properties",
    "  src/main/webapp/  css/ js/ WEB-INF/",
    "  src/test/java/    47 JUnit tests",
], size=12)
text(s, 0.6, 6.15, 6, 0.5, "Run it:  mvn jetty:run  then open localhost:8080", size=14, color=MUTED, bold=True)

# ---------------------------------------------------------------- 6. OOP
s = new_slide("OOP: inheritance, polymorphism, interfaces", "Core Java")
label_box(s, 1.5, 1.6, 3.2, 0.75, "Activity  (abstract)", "estimateCalories() uses met()", fill=BRAND, line=None, tcolor=WHITE,
          scolor=RGBColor(0xDD, 0xDF, 0xFF), tsize=14)
label_box(s, 0.5, 2.95, 3.0, 0.7, "DistanceActivity", "abstract: speed, pace", fill=BRAND_SOFT, line=None, tcolor=BRAND, tsize=13)
for i, n in enumerate(["Strength", "HIIT", "Yoga"]):
    label_box(s, 3.7 + i * 1.07, 2.95, 1.0, 0.7, n, fill=WHITE, tsize=12)
    arrow(s, 3.2, 2.35, 4.2 + i * 1.07, 2.95, color=MUTED, head=False)
arrow(s, 3.2, 2.35, 2.0, 2.95, color=MUTED, head=False)
for i, n in enumerate(["Run", "Walk", "Cycle", "Swim"]):
    label_box(s, 0.35 + i * 1.12, 4.25, 1.02, 0.7, n, fill=WHITE, tsize=12)
    arrow(s, 2.0, 3.65, 0.86 + i * 1.12, 4.25, color=MUTED, head=False)
text(s, 0.5, 5.15, 6.6, 0.4, "Polymorphism: callers hold an Activity and never test its type.", size=12, color=MUTED)
code(s, 0.5, 5.6, 6.6, 1.3, [
    "a.applyCalories(user.getWeightKg());",
    "// RunActivity.met()    -> speed based",
    "// YogaActivity.met()   -> 1.8 + 0.5*effort",
], size=12)
box(s, 7.5, 1.6, 5.25, 5.3, fill=WHITE, line=LINE, radius=0.04)
text(s, 7.75, 1.7, 4.8, 0.4, "Interfaces and abstractions", size=15, bold=True, color=BRAND)
bullets(s, 7.75, 2.15, 4.8, [
    ("Entity", "common id contract"), ("Trackable", "goals work on any activity"),
    ("Dao<T,ID>", "generic CRUD contract"), ("RowMapper<T>", "ResultSet to object"),
    ("GoogleTokenVerifier", "real and fake versions"), ("DataChangeListener", "observer"),
], size=12, gap=4)
text(s, 7.75, 4.55, 4.8, 0.4, "Exception hierarchy", size=15, bold=True, color=BRAND)
code(s, 7.75, 5.0, 4.75, 1.75, [
    "KineticException        (checked)",
    "  ValidationException  AuthException",
    "  NotFoundException",
    "DataAccessException     (unchecked)",
    "  wraps SQLException",
], size=11)

# ---------------------------------------------------------------- 7. Collections and generics
s = new_slide("Collections and generics", "Core Java")
rows = [
    ("Dao<T extends Entity, ID>", "Generic contract implemented by every DAO"),
    ("LruCache<K, V>", "Own generic cache: access-ordered LinkedHashMap + ReentrantLock"),
    ("Collection<? extends Trackable>", "Wildcard lets GoalMetric score any list of activities"),
    ("EnumMap, TreeMap, TreeSet", "Minutes by type, minutes by day, streak calculation"),
    ("ConcurrentHashMap, key sets", "Login throttle and online-users set"),
    ("Optional, records, streams", "findById results, Totals / Dashboard, groupingBy + toMap"),
    ("Comparable<Activity>", "Natural ordering by date, used by Collections.sort"),
]
for i, (a, b) in enumerate(rows):
    y = 1.65 + i * 0.62
    box(s, 0.6, y, 4.2, 0.52, fill=BRAND_SOFT, line=None, radius=0.2)
    text(s, 0.75, y, 4.0, 0.52, a, size=13, bold=True, color=BRAND, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 5.0, y, 3.4, 0.52, b, size=12, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
code(s, 8.55, 1.65, 4.2, 3.6, [
    "Map<LocalDate, List<Activity>> byWeek =",
    "  window.stream().collect(",
    "    groupingBy(a -> weekStart(a.date())));",
    "",
    "Map<LocalDate, Integer> minutesByDay =",
    "  window.stream().collect(toMap(",
    "    Activity::date,",
    "    Activity::durationMinutes,",
    "    Integer::sum, TreeMap::new));",
    "",
    "enum GoalMetric { WORKOUTS {",
    "  double measure(Collection<",
    "      ? extends Trackable> items)",
    "  { return items.size(); } }, ...",
], size=10.5)
text(s, 0.6, 6.15, 7.8, 0.6, "Streaks use a TreeSet of active days and walk backwards from today.", size=13, color=MUTED)

# ---------------------------------------------------------------- 8. Multithreading
s = new_slide("Multithreading and synchronization", "Core Java")
for i in range(3):
    label_box(s, 0.6, 1.75 + i * 0.85, 2.3, 0.7, f"Request thread {i + 1}", "log(event)", fill=WHITE, tsize=13)
    arrow(s, 2.92, 2.1 + i * 0.85, 4.1, 2.95, color=MUTED)
label_box(s, 4.1, 2.45, 2.4, 1.0, "BlockingQueue", "LinkedBlockingQueue(10,000)", fill=BRAND_SOFT, line=None, tcolor=BRAND, tsize=14)
arrow(s, 6.52, 2.95, 7.3, 2.95, color=MUTED, width=2)
label_box(s, 7.3, 2.45, 2.4, 1.0, "kinetic-audit", "daemon worker thread", fill=BRAND, line=None, tcolor=WHITE,
          scolor=RGBColor(0xDD, 0xDF, 0xFF), tsize=14)
arrow(s, 9.72, 2.95, 10.4, 2.95, color=MUTED, width=2)
label_box(s, 10.4, 2.45, 2.3, 1.0, "audit_log", "JDBC insert", fill=NAVY, line=None, tcolor=WHITE, scolor=RGBColor(0xDD, 0xDF, 0xFF), tsize=14)
text(s, 0.6, 4.4, 6, 0.4, "Where synchronization is used", size=16, bold=True, color=BRAND)
bullets(s, 0.6, 4.85, 6.2, [
    ("AuditLogger:", "producer/consumer, pages never wait for a DB write"),
    ("ConnectionPool:", "BlockingQueue of connections, 5 s wait"),
    ("LruCache:", "ReentrantLock around an access-ordered map"),
    ("Transactions:", "ThreadLocal connection per thread"),
], size=13, gap=5)
bullets(s, 6.9, 4.85, 5.8, [
    ("Login throttle:", "ConcurrentHashMap.compute is atomic"),
    ("Online users:", "ConcurrentHashMap key set"),
    ("Shutdown:", "poison pill flushes the queue"),
    ("Proof:", "ConcurrencyTest, 8 threads x 50 events = 400 rows"),
], size=13, gap=5)

# ---------------------------------------------------------------- 9. Database design
s = new_slide("Database design", "JDBC")
tables = [
    (0.6, 1.6, 3.6, "users", ["id  PK", "email  UNIQUE", "display_name, password_hash", "provider  LOCAL | GOOGLE", "google_sub  UNIQUE", "weight_kg, height_cm"]),
    (4.85, 1.6, 3.6, "activities", ["id  PK", "user_id  FK", "type, performed_on", "duration_min, distance_km", "intensity, calories", "notes"]),
    (9.1, 1.6, 3.6, "goals", ["id  PK", "user_id  FK", "metric, target", "UNIQUE(user_id, metric)"]),
    (0.6, 4.2, 3.6, "daily_logs", ["id  PK, user_id  FK", "log_date  UNIQUE per user", "steps, water_ml, sleep_min"]),
    (4.85, 4.2, 3.6, "body_metrics", ["id  PK, user_id  FK", "measured_on  UNIQUE per user", "weight_kg"]),
    (9.1, 4.2, 3.6, "audit_log", ["id  PK, user_id", "event, detail", "created_at"]),
]
for x, y, w, name, cols in tables:
    h = 0.5 + 0.28 * len(cols) + 0.15
    box(s, x, y, w, h, fill=WHITE, line=LINE, radius=0.04)
    box(s, x, y, w, 0.45, fill=BRAND, line=None, radius=0.2)
    text(s, x, y, w, 0.45, name, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, y + 0.5, w - 0.3, h - 0.55, cols, size=11, color=INK, spacing=1)
arrow(s, 4.2, 2.5, 4.85, 2.5, color=BRAND, width=2.5, head=False)
text(s, 4.2, 2.15, 0.65, 0.3, "1", size=12, bold=True, color=BRAND, align=PP_ALIGN.CENTER)
text(s, 4.2, 2.55, 0.65, 0.3, "N", size=12, bold=True, color=BRAND, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.95, 12.1, 0.4, "activities, goals, daily_logs and body_metrics reference users with ON DELETE CASCADE.", size=12, color=MUTED)
text(s, 0.6, 6.3, 12.1, 0.6, [[("DAOs: ", {"bold": True, "color": BRAND}), ("UserDao, ActivityDao, GoalDao, DailyLogDao, BodyMetricDao, AuditDao extend JdbcDao and implement Dao<T, ID>.", {})]], size=12)

# ---------------------------------------------------------------- 10. JDBC
s = new_slide("JDBC: pool, prepared statements, transactions", "JDBC")
code(s, 0.6, 1.7, 6.0, 2.2, [
    "// JdbcDao: every query is parameterised",
    "protected <T> List<T> query(String sql,",
    "        RowMapper<T> mapper, Object... params) {",
    "  return db.withConnection(c -> {",
    "    try (PreparedStatement ps =",
    "           c.prepareStatement(sql)) {",
    "      bind(ps, params);",
    "      try (ResultSet rs = ps.executeQuery()) {...}",
], size=10.5)
code(s, 0.6, 4.15, 6.0, 2.1, [
    "// GoalService: all goals save, or none do",
    "db.inTransaction(() -> {",
    "  parsed.forEach((metric, target) -> {",
    "    if (target <= 0) goals.deleteByMetric(..);",
    "    else goals.save(new Goal(...));",
    "  });",
    "  return null;",
    "});   // commit, or rollback on any exception",
], size=10.5)
bullets(s, 7.0, 1.7, 5.7, [
    ("ConnectionPool:", "fixed size, dynamic-proxy connections: close() returns to the pool"),
    ("PreparedStatement only:", "no string-built SQL, so no SQL injection"),
    ("Transactions:", "commit and rollback, nested calls join the open one"),
    ("Batch insert:", "addBatch / executeBatch seeds 16 weeks of demo data"),
    ("Generated keys:", "RETURN_GENERATED_KEYS sets the new id"),
    ("Portable:", "change db.url to use MySQL or PostgreSQL"),
    ("Tested:", "rollback test and cascade-delete test on H2"),
], size=14, gap=9)

# ---------------------------------------------------------------- 11. Servlets
s = new_slide("Servlets and web integration", "Servlets")
steps = [("Request", WHITE), ("SecurityFilter", BRAND_SOFT), ("AuthFilter", BRAND_SOFT), ("Servlet", BRAND), ("Service + DAO", CYAN), ("View", NAVY)]
for i, (t, f) in enumerate(steps):
    x = 0.6 + i * 2.1
    dark = f in (BRAND, CYAN, NAVY)
    label_box(s, x, 1.8, 1.8, 0.9, t, fill=f, line=LINE if not dark else None, tcolor=WHITE if dark else INK, tsize=13)
    if i < 5:
        arrow(s, x + 1.82, 2.25, x + 2.08, 2.25, color=MUTED, width=2)
notes = ["GET or POST", "headers, CSP, CSRF check", "redirect to /login if no session", "BaseServlet.dispatch", "business rules, JDBC", "HTML + SVG response"]
for i, n in enumerate(notes):
    text(s, 0.6 + i * 2.1, 2.78, 1.9, 0.6, n, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
box(s, 0.6, 3.6, 5.9, 2.5, fill=WHITE, line=LINE, radius=0.04)
text(s, 0.85, 3.7, 5.4, 0.4, "BaseServlet: one error policy", size=15, bold=True, color=BRAND)
bullets(s, 0.85, 4.15, 5.4, [
    "ValidationException: message shown, form re-rendered",
    "AuthException: 401 with the email kept",
    "Anything else: logged, friendly 500 page",
    "Post, redirect, get with one-time flash messages",
], size=13, gap=5)
box(s, 6.85, 3.6, 5.9, 2.5, fill=WHITE, line=LINE, radius=0.04)
text(s, 7.1, 3.7, 5.4, 0.4, "Servlet features used", size=15, bold=True, color=BRAND)
bullets(s, 7.1, 4.15, 5.4, [
    "@WebServlet, @WebFilter, @WebListener",
    "HttpSession, session id rotation, timeout",
    "ServletContextListener starts the app and DB",
    "Error pages in web.xml, CSV file download",
    "/health JSON endpoint",
], size=13, gap=5)

# ---------------------------------------------------------------- 12. Google sign-in
s = new_slide("Sign in with Google", "Authentication")
lanes = [("Browser", 1.4), ("Google", 4.6), ("Kinetic server", 8.0), ("Database", 11.3)]
for name, cx in lanes:
    label_box(s, cx - 0.9, 1.5, 1.8, 0.55, name, fill=BRAND if name == "Kinetic server" else WHITE, line=LINE if name != "Kinetic server" else None,
              tcolor=WHITE if name == "Kinetic server" else INK, tsize=13)
    arrow(s, cx, 2.1, cx, 6.7, color=LINE, width=1.5, dash=True, head=False)
msgs = [
    (1.4, 4.6, 2.55, "1  Click Continue with Google"),
    (4.6, 1.4, 3.05, "2  Signed ID token (credential)"),
    (1.4, 8.0, 3.55, "3  POST /auth/google + CSRF token"),
    (8.0, 4.6, 4.05, "4  Verify via tokeninfo"),
    (4.6, 8.0, 4.5, "5  Claims: aud, iss, exp, email"),
    (8.0, 11.3, 5.0, "6  Find, link or create user"),
    (8.0, 1.4, 5.55, "7  New session, redirect to dashboard"),
]
for x1, x2, y, label in msgs:
    arrow(s, x1, y, x2, y, color=BRAND, width=2)
    lx = min(x1, x2) + 0.1
    text(s, lx, y - 0.34, abs(x2 - x1) - 0.1, 0.3, label, size=11, color=INK, bold=True)
box(s, 0.6, 6.0, 12.15, 0.8, fill=BRAND_SOFT, line=None)
text(s, 0.85, 6.0, 11.7, 0.8, [[("Server checks:  ", {"bold": True, "color": BRAND}),
                               ("audience equals our client ID, issuer is Google, token not expired, email verified. The browser is never trusted.", {})]],
     size=13, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 13. Screens 1
s = new_slide("The application", "Screenshots")
picture(s, "landing.png", 0.6, 1.55, w=5.9)
text(s, 0.6, 5.45, 5.9, 0.3, "Landing page", size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
picture(s, "login.png", 6.85, 1.55, w=5.9)
text(s, 6.85, 5.45, 5.9, 0.3, "Sign in with Google or email", size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.6, 6.0, 12.1, 0.6, "Public pages share the same design system: indigo brand colour, rounded cards, system fonts, light and dark themes.",
     size=14, color=MUTED, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 14. Screens 2
s = new_slide("Dashboard and activity log", "Screenshots")
picture(s, "dashboard.png", 0.6, 1.5, h=4.1)
picture(s, "activities.png", 6.85, 1.5, h=4.1)
text(s, 0.6, 5.85, 5.9, 0.3, "Dashboard: totals, 12-week chart, goal rings, breakdown", size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
text(s, 6.85, 5.85, 5.9, 0.3, "Activities: log form, filters, history, CSV export", size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 15. Screens 3
s = new_slide("Health, progress and mobile", "Screenshots")
picture(s, "progress.png", 0.6, 1.5, h=3.9)
picture(s, "vitals.png", 5.75, 1.5, h=4.7)
picture(s, "mobile.png", 10.3, 1.5, h=4.7)
text(s, 0.6, 5.55, 4.7, 0.3, "Progress: calendar heatmap, records", size=12, color=MUTED, bold=True)
text(s, 5.5, 6.35, 4.0, 0.3, "Health: rings, steps, water, sleep, weight", size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
text(s, 10.1, 6.35, 2.5, 0.3, "Phone layout", size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 16. Security and testing
s = new_slide("Security and testing", "Quality")
box(s, 0.6, 1.6, 5.9, 4.3, fill=WHITE, line=LINE, radius=0.04)
text(s, 0.85, 1.7, 5.4, 0.4, "Security", size=17, bold=True, color=BRAND)
bullets(s, 0.85, 2.2, 5.4, [
    "PBKDF2 hashing with per-user salt",
    "PreparedStatement everywhere (no SQL injection)",
    "HTML escaping on all output (no XSS)",
    "CSRF token on every POST",
    "Content-Security-Policy and security headers",
    "5 failed logins per 10 minutes locks the email",
    "Users can only touch their own data",
    "Server-side Google token verification",
], size=13, gap=5)
box(s, 6.85, 1.6, 5.9, 4.3, fill=WHITE, line=LINE, radius=0.04)
text(s, 7.1, 1.7, 5.4, 0.4, "Testing: 47 JUnit tests, all passing", size=17, bold=True, color=BRAND)
bullets(s, 7.1, 2.2, 5.4, [
    "Polymorphic calorie maths and summaries",
    "Streak logic, goal scoring",
    "Registration, login, lockout",
    "Google: create, link, reject forged token",
    "Google claim checks (audience, issuer, expiry)",
    "Workout, goal, health and edit flows on in-memory JDBC",
    "Transaction rollback, cascade delete",
    "Thread safety: pool, audit logger, cache",
], size=13, gap=5)

# ---------------------------------------------------------------- 17. Rubric
s = new_slide("How the project meets the marking rubric", "Coverage")
rubric = [
    ("Problem understanding and design", "Layered architecture, ER design, clear user problem and solution"),
    ("Core Java concepts", "Inheritance, polymorphism, interfaces, exceptions, collections, generics, threads"),
    ("Database integration (JDBC)", "Own pool, PreparedStatement, transactions, batch, 6 DAO classes"),
    ("Servlets and web integration", "15 servlets, 2 filters, listener, sessions, CSRF, Google login, PWA"),
    ("Code quality", "Modular packages, comments, central error policy, 47 tests"),
]
tbl = s.shapes.add_table(len(rubric) + 1, 2, Inches(0.6), Inches(1.6), Inches(12.15), Inches(4.6)).table
tbl.rows[0].height = Inches(0.55)
for _r in range(1, len(rubric) + 1):
    tbl.rows[_r].height = Inches(0.82)
tbl.columns[0].width = Inches(4.2)
tbl.columns[1].width = Inches(7.95)
for ci, head in enumerate(["Rubric area", "Evidence in Kinetic"]):
    c = tbl.cell(0, ci)
    c.text = head
    c.fill.solid()
    c.fill.fore_color.rgb = BRAND
    p = c.text_frame.paragraphs[0]
    p.runs[0].font.bold = True
    p.runs[0].font.size = Pt(14)
    p.runs[0].font.color.rgb = WHITE
    p.runs[0].font.name = FONT
for ri, (a, b) in enumerate(rubric, start=1):
    for ci, val in enumerate((a, b)):
        c = tbl.cell(ri, ci)
        c.text = val
        c.fill.solid()
        c.fill.fore_color.rgb = WHITE if ri % 2 else BRAND_SOFT
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = c.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(14)
        r.font.bold = ci == 0
        r.font.color.rgb = INK
        r.font.name = FONT

# ---------------------------------------------------------------- 18. Thank you
s = prs.slides.add_slide(BLANK)
page["n"] += 1
set_bg(s, NAVY)
box(s, 8.3, -1.5, 7, 7, fill=RGBColor(0x1F, 0x1B, 0x6B), line=None, shape=MSO_SHAPE.OVAL)
text(s, 0.9, 1.6, 8, 1.2, "Thank you", size=60, color=WHITE, bold=True)
text(s, 0.9, 2.9, 8, 0.8, "Kinetic: move with intent.", size=24, color=RGBColor(0xC9, 0xCD, 0xF5))
box(s, 0.9, 4.3, 7.6, 1.5, fill=RGBColor(0x17, 0x1A, 0x44), line=None, radius=0.06)
text(s, 1.15, 4.4, 7.2, 1.3, [
    [("Run  ", {"bold": True, "color": LIME}), ("mvn jetty:run  then  localhost:8080", {"color": WHITE})],
    [("Demo  ", {"bold": True, "color": LIME}), ("Try the live demo button, or demo@kinetic.app", {"color": WHITE})],
], size=17, spacing=10, anchor=MSO_ANCHOR.MIDDLE)
picture(s, "mobile.png", 10.4, 1.3, h=5.2, frame=False)

out = os.path.join(HERE, "Kinetic-Presentation.pptx")
prs.save(out)
print("saved", out, "slides:", page["n"])
